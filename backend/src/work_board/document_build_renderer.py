"""Fixed CPU-only rendering; canonical filesystem publication is not owned here."""
from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO
from xml.sax.saxutils import escape
from .document_build_contracts import DocumentBuildSpec, MAX_FILE_BYTES, DOCX_MEDIA, XLSX_MEDIA, coordinate
from .document_build_formula import calculate


@dataclass(frozen=True)
class RenderResult:
    editable: bytes
    editable_media: str
    editable_extension: str
    pdf: bytes | None
    warnings: list[str]
    source_refs: list[str]


def display(value):
    return '' if value is None else 'TRUE' if value is True else 'FALSE' if value is False else str(value)


def content(spec, values):
    """One semantic view used by the direct PDF renderer."""
    result = [('heading', spec.title)]
    labels = {citation.source_ref: citation.label for citation in spec.citations}
    for section in spec.sections:
        result.append(('heading', section.heading))
        result.extend(('paragraph', paragraph) for paragraph in section.paragraphs)
        if section.citation_refs:
            result.append(('paragraph', '; '.join(labels[ref] for ref in section.citation_refs)))
    if spec.kind == 'table_workbook':
        workbook = spec.tables[0]
        for sheet in workbook.sheet_names:
            result.append(('heading', sheet))
            # Sparse cells are represented by exact coordinates, without huge
            # implicit blank grids or silent omitted values.
            rows = [['Cell', 'Value']]
            for (name, row, column), value in sorted(values.items()):
                if name == sheet:
                    letters = ''
                    index = column + 1
                    while index:
                        index, rem = divmod(index - 1, 26)
                        letters = chr(65 + rem) + letters
                    rows.append([f'{letters}{row+1}', display(value)])
            result.append(('table', rows))
    else:
        for table in spec.tables:
            result.append(('heading', table.title))
            result.append(('table', [table.columns, *[[display(value) for value in row] for row in table.rows]]))
            if table.citation_refs:
                result.append(('paragraph', '; '.join(labels[ref] for ref in table.citation_refs)))
    if spec.citations:
        result.append(('heading', 'Sources'))
        result.extend(('paragraph', citation.label + ': ' + citation.source_ref) for citation in spec.citations)
    return result


def docx(spec, semantic):
    from docx import Document
    from docx.shared import Inches, Pt
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.75)
        section.left_margin = section.right_margin = Inches(.75)
    document.styles['Normal'].font.size = Pt(10 if spec.style_preset == 'compact' else 11)
    for kind, value in semantic:
        if kind == 'heading':
            document.add_heading(value, level=1)
        elif kind == 'paragraph':
            document.add_paragraph(value)
        else:
            table = document.add_table(rows=len(value), cols=len(value[0]))
            table.style = 'Table Grid'
            for row, cells in zip(table.rows, value):
                for cell, text in zip(row.cells, cells):
                    cell.text = text
    target = BytesIO()
    document.save(target)
    return target.getvalue()


def xlsx(spec, values):
    import xlsxwriter
    target = BytesIO()
    workbook = xlsxwriter.Workbook(target, {'in_memory': True, 'strings_to_formulas': False, 'strings_to_urls': False})
    spreadsheet = spec.tables[0]
    sheets = {name: workbook.add_worksheet(name) for name in spreadsheet.sheet_names}
    overview = workbook.add_worksheet('Overview')
    row = 0
    for kind, value in content(spec, {})[:1 + sum(1 + len(s.paragraphs) + bool(s.citation_refs) for s in spec.sections)]:
        overview.write_string(row, 0, value)
        row += 1
    for citation in spec.citations:
        overview.write_string(row, 0, citation.label)
        overview.write_string(row, 1, citation.source_ref)
        row += 1
    presets = {
        'plain': {}, 'integer': {'num_format': '0'}, 'decimal': {'num_format': '0.00'},
        'percent': {'num_format': '0.00%'}, 'currency': {'num_format': '#,##0.00'},
        'header': {'bold': True, 'bg_color': '#E5E7EB'},
    }
    formats = {name: workbook.add_format(options) for name, options in presets.items()}
    assigned = {(entry.sheet, *coordinate(entry.cell)): formats[entry.preset] for entry in spreadsheet.formats}
    occupied = {(entry.sheet, *coordinate(entry.cell)) for entry in [*spreadsheet.cells, *spreadsheet.formulas]}
    for (name, row, col), fmt in assigned.items():
        if (name, row, col) not in occupied:
            sheets[name].write_blank(row, col, None, fmt)
    for entry in spreadsheet.cells:
        row, col = coordinate(entry.cell)
        sheet, value = sheets[entry.sheet], entry.value
        fmt = assigned.get((entry.sheet, row, col))
        if type(value) is str:
            result = sheet.write_string(row, col, value, fmt)
        elif type(value) is bool:
            result = sheet.write_boolean(row, col, value, fmt)
        elif value is None:
            result = sheet.write_blank(row, col, None, fmt)
        else:
            result = sheet.write_number(row, col, value, fmt)
        if result != 0:
            raise ValueError('document_xlsx_write_failed')
    for entry in spreadsheet.formulas:
        row, col = coordinate(entry.cell)
        value = values[(entry.sheet, row, col)]
        if value is None or isinstance(value, list):
            raise ValueError('document_formula_scalar_result_required')
        expression = entry.expression if entry.expression.startswith('=') else '=' + entry.expression
        if sheets[entry.sheet].write_formula(row, col, expression, assigned.get((entry.sheet, row, col)), value) != 0:
            raise ValueError('document_xlsx_write_failed')
    workbook.close()
    return target.getvalue()


def pdf(spec, semantic):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    # Built-in Helvetica cannot faithfully render arbitrary Unicode; never
    # emit replacement glyphs while claiming the same content.
    text_values = [text for kind, value in semantic for text in ([value] if kind != 'table' else [cell for row in value for cell in row])]
    if any(any(ord(char) > 255 or ord(char) < 32 and char not in '\n\t\r' for char in text) for text in text_values):
        raise ValueError('document_pdf_unsupported_glyph')
    styles = getSampleStyleSheet()
    styles['Normal'].fontSize = 9 if spec.style_preset == 'compact' else 10
    styles['Normal'].leading = 12
    story = []
    for kind, value in semantic:
        if kind == 'table':
            rows = [[Paragraph(escape(cell).replace('\n', '<br/>'), styles['Normal']) for cell in row] for row in value]
            table = LongTable(rows, colWidths=[(A4[0]-72)/len(value[0])] * len(value[0]), repeatRows=1, splitByRow=1, splitInRow=1)
            table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .3, colors.grey), ('VALIGN', (0,0), (-1,-1), 'TOP'), ('BACKGROUND', (0,0), (-1,0), colors.lightgrey)]))
            story.append(table)
        else:
            story.append(Paragraph(escape(value).replace('\n', '<br/>'), styles['Heading1'] if kind == 'heading' else styles['Normal']))
        story.append(Spacer(1, 6))
    target = BytesIO()
    def page(canvas, document):
        if document.page > 100:
            raise ValueError('document_pdf_page_bound')
    SimpleDocTemplate(target, pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36).build(story, onFirstPage=page, onLaterPages=page)
    return target.getvalue()


def render(spec: DocumentBuildSpec, selected_view=None) -> RenderResult:
    if type(spec) is not DocumentBuildSpec:
        raise TypeError('document_validated_spec_required')
    # Revalidate immutable DTO arrays at execution boundary; frozen models do
    # not make their nested Python lists immutable.
    spec = DocumentBuildSpec.model_validate(spec.model_dump())
    values = calculate(spec.tables[0]) if spec.kind == 'table_workbook' else {}
    semantic = content(spec, values)
    editable = xlsx(spec, values) if spec.kind == 'table_workbook' else docx(spec, semantic)
    if not 0 < len(editable) <= MAX_FILE_BYTES:
        raise ValueError('document_editable_byte_bound')
    warnings = []
    try:
        rendered = pdf(spec, semantic)
        if not 0 < len(rendered) <= MAX_FILE_BYTES:
            raise ValueError('document_pdf_byte_bound')
    except Exception as exc:
        rendered = None
        warnings.append(str(exc) if str(exc) in ('document_pdf_unsupported_glyph', 'document_pdf_page_bound', 'document_pdf_byte_bound') else 'document_pdf_render_unavailable')
    return RenderResult(editable, XLSX_MEDIA if spec.kind == 'table_workbook' else DOCX_MEDIA, 'xlsx' if spec.kind == 'table_workbook' else 'docx', rendered, warnings, [citation.source_ref for citation in spec.citations])
