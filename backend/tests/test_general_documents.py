"""Literal local formats, confined process and private owner-selected readback."""
import asyncio
import hashlib
import io
import json
import os
import sys
import zipfile

import httpx
import pytest
from fastapi import FastAPI
from tests.test_inference_accounting import accounting_db
from tests.document_compare_fixtures import invoice_pdf
from src.work_board.documents import DocumentReadInput, DocumentService
from src.work_board.document_read_parser import extract, DocumentReadError, office_package


def fixture_bytes(fmt):
    if fmt == "pdf":
        return invoice_pdf(["Selected owner document", "A literal source statement"])
    if fmt == "csv":
        return b'Name,Count,Literal\r\nAlpha,2,=SUM(B2:B3)\r\n'
    output = io.BytesIO()
    if fmt == "docx":
        from docx import Document
        document = Document(); document.add_paragraph("Selected owner paragraph")
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Private table"
        document.save(output)
    else:
        from openpyxl import Workbook
        workbook = Workbook(); workbook.active.title = "Inputs"
        workbook.active.append(["Name", "Count", "Literal"])
        workbook.active.append(["Alpha", 2, "=SUM(B2:B3)"])
        workbook.create_sheet("Skipped").append(["not selected"])
        workbook.save(output); workbook.close()
    return output.getvalue()


def read_request(fmt, identifier="00000000-0000-0000-0000-000000000000"):
    return DocumentReadInput(artifact_ref="document-source:"+identifier, format=fmt,
        selection={"sheets": ["Inputs"], "pages": []} if fmt == "xlsx" else {"pages": [], "sheets": []})


async def test_upload_profile_failed_child_is_reaped_and_restart_proves_real_filesystem(accounting_db, monkeypatch):
    from src.work_board import document_pairs as sources
    from src.work_board.repository import BoardError
    original_spawn = asyncio.create_subprocess_exec
    children = []
    async def silent_child(*args, **kwargs):
        process = await original_spawn(sys.executable, "-I", "-c", "import signal,time; signal.alarm(2); time.sleep(5)", **kwargs)
        children.append(process)
        return process
    monkeypatch.setattr(asyncio, "create_subprocess_exec", silent_child)
    service = DocumentService()
    try:
        await service.start()
        assert service._started and service._upload_profile is None
        assert len(children) == 1 and children[0].returncode is not None
        with pytest.raises(BoardError) as absent: sources.validate_upload_profile(service._upload_profile)
        assert absent.value.code == "document_upload_profile_unproved"
        monkeypatch.setattr(asyncio, "create_subprocess_exec", original_spawn)
        await service.start()
        proven = service._upload_profile
        assert proven["cross_process"] and proven["positive_wait"]
        sources.validate_upload_profile(proven)
        await service.stop()
        with pytest.raises(BoardError): sources.validate_upload_profile(proven)
    finally:
        await service.stop()


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx", "csv"])
async def test_real_format_confined_process_exact_citations(fmt):
    service = DocumentService(); await service.start()
    raw = fixture_bytes(fmt); request = read_request(fmt)
    result = await service.parse(raw, request)
    await service.stop()
    assert result["status"] == "succeeded", result
    assert result["cleanup"] == "wait_reaped" and result["provider_contacts"] == 0
    evidence = result["evidence"]
    assert evidence["source_digest"] == hashlib.sha256(raw).hexdigest()
    assert evidence["no_learning"] is True
    citations = [section["source_ref"] for section in evidence["sections"]]
    assert all(ref.startswith(request.artifact_ref+"#") for ref in citations)
    assert {"pdf": "page=1", "docx": "paragraph=1", "xlsx": "sheet=Inputs&row=1", "csv": "sheet=CSV&row=1"}[fmt] in citations[0]
    if fmt in {"xlsx", "csv"}:
        cells = [cell for section in evidence["sections"] for cell in section["table_cells"]]
        formula = next(cell for cell in cells if cell["formula"])
        assert formula["formula"] == "=SUM(B2:B3)" and formula["cached_value"] is None
        assert "cell=C2" in formula["source_ref"]
    if fmt == "xlsx":
        assert "Skipped XLSX sheets: Skipped" in evidence["warnings"]


@pytest.mark.parametrize("name,raw,reason", [
    ("../outside.xml", b"x", "document_zip_path_unsupported"),
    ("xl/vbaProject.bin", b"macro", "document_macros_or_external_links_unsupported"),
    ("xl/externalLinks/link.xml", b"link", "document_macros_or_external_links_unsupported"),
    ("word/document.xml", b"x"*100000, "document_zip_expansion_or_protection_unsupported"),
    ("word/_rels/document.xml.rels", b'<Relationship TargetMode="External" Target="http://fixture.invalid"/>', "document_external_relationship_unsupported"),
    ("word/document.xml", b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>', "document_xml_entities_unsupported"),
])
def test_office_preflight_bomb_traversal_and_active_content(name, raw, reason):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr(name, raw)
    with pytest.raises(DocumentReadError, match=reason):
        office_package(buffer.getvalue())


def test_finite_selection_output_and_cells():
    with pytest.raises(ValueError):
        DocumentReadInput(artifact_ref="/etc/passwd", format="pdf")
    with pytest.raises(ValueError):
        DocumentReadInput(artifact_ref=read_request("pdf").artifact_ref, format="csv", selection={"pages": [1]})
    request = read_request("csv").model_dump(); request["page_sheet_limits"]["max_cells"] = 1
    with pytest.raises(DocumentReadError, match="document_populated_cell_limit_exceeded"):
        extract(b"a,b", request)
    with pytest.raises(DocumentReadError, match="document_output_size_exceeded"):
        extract(b"a"*(1024*1024), read_request("csv").model_dump())


async def test_timeout_cleanup_and_inactive_state():
    service = DocumentService()
    with pytest.raises(Exception, match="Start the managed document service"):
        await service.parse(b"a", read_request("csv"))
    await service.start()
    result = await service.parse(fixture_bytes("csv"), read_request("csv"), timeout=.001)
    assert result["status"] == "blocked" and result["cleanup"] == "wait_reaped"
    assert not service._processes and not service._capacity.locked()
    await service.stop()


async def test_parser_scanned_protected_malformed_and_crash(monkeypatch):
    from pypdf import PdfWriter
    service = DocumentService(); await service.start()
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    output = io.BytesIO(); writer.write(output)
    result = await service.parse(output.getvalue(), read_request("pdf"))
    assert result["reason"] == "document_pdf_scanned_or_empty_unsupported"
    writer.encrypt("private-password"); output = io.BytesIO(); writer.write(output)
    result = await service.parse(output.getvalue(), read_request("pdf"))
    assert result["reason"] == "document_pdf_protected_unsupported"
    result = await service.parse(b"not a PDF", read_request("pdf"))
    assert result["reason"] == "document_malformed_or_unsupported"
    await service.stop()


async def test_spreadsheet_literal_formula_and_cached_value_are_separate():
    original = fixture_bytes("xlsx")
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                raw = raw.replace(b"<f>SUM(B2:B3)</f><v></v>", b"<f>SUM(B2:B3)</f><v>4</v>")
            target.writestr(item, raw)
    service = DocumentService(); await service.start()
    result = await service.parse(output.getvalue(), read_request("xlsx"))
    await service.stop()
    cells = [cell for section in result["evidence"]["sections"] for cell in section["table_cells"]]
    formula = next(cell for cell in cells if cell["formula"])
    assert formula["text"] == formula["formula"] == "=SUM(B2:B3)" and formula["cached_value"] == "4"
    assert "cached value freshness is unknown" in result["evidence"]["warnings"][-1]


def formula_fixture(formula_xml):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(fixture_bytes("xlsx"))) as source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            raw = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                original = b"<f>SUM(B2:B3)</f><v></v>"
                assert original in raw
                raw = raw.replace(original, formula_xml+b"<v>4</v>")
            target.writestr(item, raw)
    return output.getvalue()


async def test_array_formula_exact_literal_cache_and_deterministic_child_readback():
    raw = formula_fixture(b'<f t="array" ref="C2:C3">SUM(B2:B3)</f>')
    service = DocumentService(); await service.start()
    try:
        results = [await service.parse(raw, read_request("xlsx")) for _ in range(2)]
    finally:
        await service.stop()
    assert results[0] == results[1]
    assert results[0]["status"] == "succeeded", results[0]
    cells = [cell for section in results[0]["evidence"]["sections"] for cell in section["table_cells"]]
    formula = next(cell for cell in cells if cell["formula"])
    assert formula["text"] == formula["formula"] == "=SUM(B2:B3)"
    assert formula["cached_value"] == "4" and "cell=C2" in formula["source_ref"]
    assert " object at " not in json.dumps(results)


@pytest.mark.parametrize("formula_xml,reason", [
    (b'<f t="dataTable" ref="C2:C3" r1="B2"/>', "document_spreadsheet_data_table_formula_unsupported"),
    (b'<f t="array" ref="C2:C3"/>', "document_spreadsheet_formula_literal_unsupported"),
])
async def test_nonliteral_formula_objects_have_explicit_unsupported_state(formula_xml, reason):
    service = DocumentService(); await service.start()
    try:
        result = await service.parse(formula_fixture(formula_xml), read_request("xlsx"))
    finally:
        await service.stop()
    assert result["status"] == "blocked" and result["reason"] == reason, result
    assert result["cleanup"] == "wait_reaped" and result["provider_contacts"] == 0
    assert "evidence" not in result


def test_unknown_formula_objects_never_serialize_repr():
    from src.work_board.document_read_parser import spreadsheet_formula_literal
    class UnsupportedFormula:
        def __str__(self):
            raise AssertionError("formula object repr was consumed")
    with pytest.raises(DocumentReadError, match="document_spreadsheet_formula_literal_unsupported"):
        spreadsheet_formula_literal(UnsupportedFormula())


@pytest.mark.parametrize("fmt", ["pdf", "docx", "xlsx", "csv", "restart", "cancel",
    "upload_cancel", "upload_restart", "upload_inode", "upload_symlink", "parent_open",
    "root_revoke", "root_expire", "goal_pause", "goal_revision", "original_deadline", "late_cancel",
    "profile_missing", "profile_moved", "profile_stop", "upload_orphan"])
async def test_authenticated_private_source_upload_seal_read_delete(accounting_db, monkeypatch, fmt):
    from config.settings import settings
    from src.api import auth, documents
    from src.api.router import api_router
    from src.auth.middleware import OperatorAuthMiddleware
    from src.vault import crypto
    from src.db.models import WorkBoardInputArtifact
    mode = fmt
    if fmt not in {"pdf", "docx", "xlsx", "csv"}:
        fmt = "csv"
    root, _engine, factory = accounting_db
    monkeypatch.setattr(documents, "get_session", factory.accounting_sessions)
    monkeypatch.setattr(crypto, "_fernet", None)
    monkeypatch.setattr(settings, "vault_encryption_key", "")
    monkeypatch.setattr(settings, "operator_auth_allow_unauthenticated_tests", False)
    monkeypatch.setattr(settings, "operator_auth_secret", "isolated-document-test")
    monkeypatch.setattr(settings, "operator_auth_secret_hash", "")
    monkeypatch.setattr(settings, "operator_auth_allowed_hosts", "test,localhost,127.0.0.1")
    monkeypatch.setattr(settings, "operator_auth_allowed_origins", "http://localhost:3001")
    monkeypatch.setattr(settings, "operator_auth_cookie_secure", False)
    # Any accidental provider HTTP entry fails before transport; child sockets
    # are independently denied by its kernel self-check.
    contacts = []
    async def deny_provider(*args, **kwargs):
        contacts.append(True); raise AssertionError("document extraction contacted inference")
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", deny_provider)
    import litellm
    monkeypatch.setattr(litellm, "acompletion", deny_provider)
    def deny_provider_sync(*args, **kwargs):
        contacts.append(True); raise AssertionError("document extraction contacted inference")
    monkeypatch.setattr(litellm, "completion", deny_provider_sync)
    auth._reset_login_throttle_for_tests()
    app = FastAPI(); app.add_middleware(OperatorAuthMiddleware)
    app.include_router(api_router)
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test",
            headers={"origin": "http://localhost:3001"}) as client:
            assert (await client.post("/api/documents/read", json=read_request(fmt).model_dump())).status_code == 401
            assert (await client.post("/api/auth/login", json={"password": "isolated-document-test"})).status_code == 200
            goal = await client.post("/api/goals", json={"title": "Read selected private document",
                "admission_budget": {"reviewed_grant": True, "grant_id": "document-local-test",
                    "max_outstanding_jobs": 1, "max_attempts": 2, "max_runtime_seconds": 70}})
            assert goal.status_code == 200, goal.text
            raw = fixture_bytes(fmt)
            reserved = await client.post("/api/documents/sources", json={"format": fmt,
                "source": {"size_bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()},
                "goal_id": goal.json()["id"], "goal_revision": 1, "idempotency_key": "selected", "no_learning": True})
            assert reserved.status_code == 200, reserved.text
            state = reserved.json(); identifier = state["artifact_id"]
            if mode.startswith("profile_"):
                service = app.state.document_service
                old_directory = None
                if mode == "profile_missing": service._upload_profile = None
                elif mode == "profile_stop": await service.stop()
                else:
                    source_directory = root/"artifacts/work-board/document-sources"
                    old_directory = source_directory.with_name("original-source-directory")
                    source_directory.rename(old_directory); source_directory.mkdir(mode=0o700)
                try:
                    refused = await client.put(f"/api/documents/sources/{identifier}/content", params={"expected_revision": state["revision"]},
                        content=raw, headers={"content-type": "application/octet-stream"})
                    assert refused.status_code == 503 and refused.json()["detail"]["code"] == "document_upload_profile_unproved"
                    retained_index = await client.get("/api/documents/sources")
                    assert retained_index.status_code == 200 and retained_index.json()["upload_readiness"] == "blocked"
                    assert retained_index.json()["sources"][0]["artifact_id"] == identifier
                    async with factory.accounting_sessions() as db:
                        unchanged = await db.get(WorkBoardInputArtifact, identifier)
                        assert unchanged.revision == state["revision"] and json.loads(unchanged.document_metadata_json)["live_writer"] is None
                    assert contacts == []
                finally:
                    if old_directory is not None:
                        source_directory.rmdir(); old_directory.rename(source_directory)
                return
            if mode == "upload_orphan":
                from src.work_board import document_pairs as sources
                from src.work_board.contracts import WorkBoardOwner
                async with factory.accounting_sessions() as db:
                    source_row = await db.get(WorkBoardInputArtifact, identifier)
                    orphan_owner = WorkBoardOwner(principal_id=source_row.owner_principal_id, session_id=source_row.owner_session_id)
                lease_inode = None
                for _ in range(3):
                    async with factory() as db:
                        async def crash_commit(): raise RuntimeError("isolated owner failure before upload reservation commit")
                        monkeypatch.setattr(db, "commit", crash_commit)
                        with pytest.raises(RuntimeError):
                            await sources.acquire_upload(db, orphan_owner, identifier, state["revision"], "source", capability=sources.SOURCE_CAPABILITY,
                                upload_profile=app.state.document_service._upload_profile)
                        await db.rollback()
                    leases = list(root.glob(f"artifacts/work-board/document-sources/{identifier}/*.upload-lock"))
                    assert len(leases) == 1
                    if lease_inode is None: lease_inode = leases[0].stat().st_ino
                    assert leases[0].stat().st_ino == lease_inode
                    async with factory.accounting_sessions() as db:
                        unchanged = await db.get(WorkBoardInputArtifact, identifier)
                        assert unchanged.revision == state["revision"] and json.loads(unchanged.document_metadata_json)["live_writer"] is None
            if mode.startswith("upload_") and mode != "upload_orphan":
                from src.work_board import document_pairs as sources
                from src.work_board.contracts import WorkBoardOwner
                entered, closed = asyncio.Event(), asyncio.Event()
                async with factory.accounting_sessions() as db:
                    row = await db.get(WorkBoardInputArtifact, identifier)
                    owner = WorkBoardOwner(principal_id=row.owner_principal_id, session_id=row.owner_session_id)
                if mode == "upload_cancel":
                    async def interrupted_stream():
                        try:
                            yield raw[:1]
                            entered.set()
                            await asyncio.Event().wait()
                        finally:
                            closed.set()
                    async def active_upload():
                        async with factory.accounting_sessions() as db:
                            await sources.upload(db, owner, identifier, state["revision"], "source", interrupted_stream(), capability=sources.SOURCE_CAPABILITY,
                                upload_profile=app.state.document_service._upload_profile)
                    running_upload = asyncio.create_task(active_upload())
                    await asyncio.wait_for(entered.wait(), 5)
                    live = (await client.get(f"/api/documents/sources/{identifier}")).json()
                    assert live["writer_kind"] == "upload"
                    refused = await client.post(f"/api/documents/sources/{identifier}/reconcile-upload", params={"expected_revision": live["revision"]})
                    assert refused.status_code == 409 and refused.json()["detail"]["code"] == "document_upload_quiescence_unknown"
                    running_upload.cancel()
                    with pytest.raises(asyncio.CancelledError): await asyncio.wait_for(running_upload, 5)
                    assert closed.is_set()
                    settled = (await client.get(f"/api/documents/sources/{identifier}")).json()
                    assert settled["cleanup"] == "quiescent" and settled["state"] == "cleanup_required"
                else:
                    async with factory.accounting_sessions() as db:
                        upload_row, upload_value, _token, lease_fd = await sources.acquire_upload(db, owner, identifier, state["revision"], "source", capability=sources.SOURCE_CAPABILITY,
                            upload_profile=app.state.document_service._upload_profile)
                    # Transfer the original open file description to a real
                    # process; its death releases the exact kernel-held lease.
                    holder = await asyncio.create_subprocess_exec(sys.executable, "-I", "-c",
                        "import sys; print('holding',flush=True); sys.stdin.buffer.read()",
                        stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
                        stderr=asyncio.subprocess.DEVNULL, pass_fds=(lease_fd,), close_fds=True, env={})
                    os.close(lease_fd)
                    try:
                        assert await asyncio.wait_for(holder.stdout.readline(), 5) == b"holding\n"
                        live = (await client.get(f"/api/documents/sources/{identifier}")).json()
                        refused = await client.post(f"/api/documents/sources/{identifier}/reconcile-upload", params={"expected_revision": live["revision"]})
                        assert refused.status_code == 409
                        refused_delete = await client.delete(f"/api/documents/sources/{identifier}", params={"expected_revision": live["revision"]})
                        assert refused_delete.status_code == 409
                    finally:
                        if holder.returncode is None: holder.kill()
                        await asyncio.wait_for(holder.wait(), 5)
                    lease_path = sources.source_path(upload_row, upload_value, "source").parent / upload_value["upload_binding"]["file"]
                    if mode in {"upload_inode", "upload_symlink"}:
                        saved = lease_path.with_name("saved-original")
                        lease_path.rename(saved)
                        if mode == "upload_inode":
                            lease_path.write_bytes(saved.read_bytes()); lease_path.chmod(0o600)
                        else:
                            lease_path.symlink_to(saved)
                        refused = await client.post(f"/api/documents/sources/{identifier}/reconcile-upload", params={"expected_revision": live["revision"]})
                        assert refused.status_code == 409
                        async with factory.accounting_sessions() as db:
                            retained_row = await db.get(WorkBoardInputArtifact, identifier)
                            assert retained_row.document_reserved_bytes == 32*1024*1024 and sources.metadata(retained_row)["live_writer"]
                        lease_path.unlink(); saved.rename(lease_path)
                    settled_response = await client.post(f"/api/documents/sources/{identifier}/reconcile-upload", params={"expected_revision": live["revision"]})
                    assert settled_response.status_code == 200, settled_response.text
                    settled = settled_response.json()
                    assert settled["state"] == "cleanup_required" and settled["cleanup"] == "quiescent"
                async with factory.accounting_sessions() as db:
                    retained_row = await db.get(WorkBoardInputArtifact, identifier)
                    assert retained_row.document_reserved_bytes == 32*1024*1024 and not sources.metadata(retained_row)["sources"]
                assert (await client.post(f"/api/documents/sources/{identifier}/seal", params={"expected_revision": settled["revision"]})).status_code == 409
                deleted = await client.delete(f"/api/documents/sources/{identifier}", params={"expected_revision": settled["revision"]})
                assert deleted.status_code == 200 and deleted.json()["state"] == "deleted", deleted.text
                async with factory.accounting_sessions() as db:
                    assert (await db.get(WorkBoardInputArtifact, identifier)).document_reserved_bytes == 0
                assert contacts == []
                return
            for route in ("/api/documents/sources/"+identifier,):
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test", headers={"origin": "http://localhost:3001"}) as other:
                    assert (await other.get(route)).status_code == 401
                    assert (await other.post("/api/auth/login", json={"password": "isolated-document-test"})).status_code == 200
                    assert (await other.get(route)).status_code == 404
            uploaded = await client.put(f"/api/documents/sources/{identifier}/content", params={"expected_revision": state["revision"]},
                content=raw, headers={"content-type": "application/octet-stream"})
            assert uploaded.status_code == 200, uploaded.text
            state = uploaded.json()
            sealed = await client.post(f"/api/documents/sources/{identifier}/seal", params={"expected_revision": state["revision"]})
            assert sealed.status_code == 200, sealed.text
            assert (await client.get(f"/api/work-board/document-pairs/{identifier}")).status_code == 404
            if mode == "csv":
                invoice = invoice_pdf()
                csv_source = b"SKU,QTY,UNIT_PRICE\nPEN-01,2,3.50\nBOOK-02,1,12.00\n"
                paired = await client.post("/api/work-board/document-pairs", json={"schema_version": 1,
                    "operation": "compare-line-totals-by-sku", "goal_id": goal.json()["id"], "goal_revision": 1,
                    "idempotency_key": "mixed-family", "no_learning": True,
                    "pdf": {"size_bytes": len(invoice), "sha256": hashlib.sha256(invoice).hexdigest()},
                    "csv": {"size_bytes": len(csv_source), "sha256": hashlib.sha256(csv_source).hexdigest()}})
                assert paired.status_code == 200, paired.text
                pair = paired.json()
                assert (await client.get(f"/api/documents/sources/{pair['artifact_id']}")).status_code == 404
                third = await client.post("/api/documents/sources", json={"format": fmt,
                    "source": {"size_bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()},
                    "goal_id": goal.json()["id"], "goal_revision": 1, "idempotency_key": "quota-overflow", "no_learning": True})
                assert third.status_code == 409 and third.json()["detail"]["code"] == "document_pair_quota_full"
                discarded = await client.post(f"/api/work-board/document-pairs/{pair['artifact_id']}/discard", json={"expected_revision": pair["revision"]})
                assert discarded.status_code == 200 and discarded.json()["quota_reserved_bytes"] == 0
            if mode in {"restart", "cancel"}:
                second = await client.post("/api/documents/sources", json={"format": fmt,
                    "source": {"size_bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()},
                    "goal_id": goal.json()["id"], "goal_revision": 1, "idempotency_key": "second-selected", "no_learning": True})
                assert second.status_code == 200, second.text
                second_id = second.json()["artifact_id"]
                second_upload = await client.put(f"/api/documents/sources/{second_id}/content", params={"expected_revision": second.json()["revision"]},
                    content=raw, headers={"content-type": "application/octet-stream"})
                second_seal = await client.post(f"/api/documents/sources/{second_id}/seal", params={"expected_revision": second_upload.json()["revision"]})
                assert second_seal.status_code == 200, second_seal.text
                old = app.state.document_service
                original_parse = old.parse
                ready_event, release_event = asyncio.Event(), asyncio.Event()
                async def paused_parse(*args, **kwargs):
                    original_ready = kwargs["on_ready"]
                    async def paused_ready(packet):
                        await original_ready(packet)
                        ready_event.set()
                        await release_event.wait()
                    kwargs["on_ready"] = paused_ready
                    return await original_parse(*args, **kwargs)
                monkeypatch.setattr(old, "parse", paused_parse)
                running = asyncio.create_task(client.post("/api/documents/read", json=read_request(fmt, identifier).model_dump()))
                await asyncio.wait_for(ready_event.wait(), 5)
                assert old._supervisors
                # A fresh service has a fresh asyncio lock but must obey the
                # persisted host slot for this actually living old child.
                replacement = DocumentService(); await replacement.start()
                app.state.document_service = replacement
                denied = await client.post("/api/documents/read", json=read_request(fmt, second_id).model_dump())
                assert denied.status_code == 409 and denied.json()["detail"]["code"] == "document_parser_capacity_held"
                original_state = (await client.get(f"/api/documents/sources/{identifier}")).json()
                refused = await client.delete(f"/api/documents/sources/{identifier}", params={"expected_revision": original_state["revision"]})
                assert refused.status_code == 409 and refused.json()["detail"]["code"] == "document_upload_quiescence_unknown"
                if mode == "cancel":
                    running.cancel()
                    with pytest.raises(asyncio.CancelledError):
                        await asyncio.wait_for(running, 5)
                else:
                    # Killing the supervisor would lose its witness; forbid it.
                    for supervisor in tuple(old._supervisors):
                        def forbidden_kill():
                            raise AssertionError("shutdown killed the positive witness owner")
                        monkeypatch.setattr(supervisor, "kill", forbidden_kill)
                    await asyncio.wait_for(old.stop(), 5)
                    release_event.set()
                    interrupted = await asyncio.wait_for(running, 5)
                    assert interrupted.json()["status"] == "blocked", interrupted.text
                assert not old._processes
                witnesses = list(root.glob(f"artifacts/work-board/document-sources/{identifier}/*.witness.json"))
                assert len(witnesses) == 1 and json.loads(witnesses[0].read_text())["wait_reaped"] is True
                # Simulate parent death before committing its already-written
                # positive wait. Missing/wrong witnesses cannot free the slot.
                async with factory.accounting_sessions() as db:
                    row = await db.get(WorkBoardInputArtifact, identifier)
                    metadata = json.loads(row.document_metadata_json)
                    metadata["live_writer"] = {"slot": "parser", "token": metadata["parser_binding"]["nonce"]}
                    row.document_metadata_json = json.dumps(metadata)
                    from src.work_board.input_artifacts import _metadata_digest
                    row.metadata_digest = _metadata_digest(row)
                original_state = (await client.get(f"/api/documents/sources/{identifier}")).json()
                original_witness = witnesses[0].read_bytes()
                wrong_witness = json.loads(original_witness); wrong_witness["parser_pid"] += 1
                witnesses[0].write_text(json.dumps(wrong_witness))
                bad_recovery = await client.post(f"/api/documents/sources/{identifier}/reconcile", params={"expected_revision": original_state["revision"]})
                assert bad_recovery.status_code == 409 and bad_recovery.json()["detail"]["code"] == "document_parser_cleanup_unknown"
                still_denied = await client.post("/api/documents/read", json=read_request(fmt, second_id).model_dump())
                assert still_denied.status_code == 409 and still_denied.json()["detail"]["code"] == "document_parser_capacity_held"
                witnesses[0].write_bytes(original_witness)
                recovered = await client.post(f"/api/documents/sources/{identifier}/reconcile", params={"expected_revision": original_state["revision"]})
                assert recovered.status_code == 200 and recovered.json()["cleanup"] == "quiescent", recovered.text
                second_result = await client.post("/api/documents/read", json=read_request(fmt, second_id).model_dump())
                assert second_result.json()["status"] == "succeeded", second_result.text
                second_state = (await client.get(f"/api/documents/sources/{second_id}")).json()
                assert (await client.delete(f"/api/documents/sources/{second_id}", params={"expected_revision": second_state["revision"]})).status_code == 200
            if mode == "parent_open":
                from src.work_board import documents as document_service
                original_open = document_service._open_input_artifact_parent
                def denied_parent(*args, **kwargs): raise OSError("isolated prelaunch descriptor failure")
                monkeypatch.setattr(document_service, "_open_input_artifact_parent", denied_parent)
                denied = await client.post("/api/documents/read", json=read_request(fmt, identifier).model_dump())
                assert denied.status_code == 409
                async with factory.accounting_sessions() as db:
                    current_row = await db.get(WorkBoardInputArtifact, identifier)
                    retained_value = json.loads(current_row.document_metadata_json)
                    assert retained_value["live_writer"] is None and retained_value.get("parser_attempts", 0) == 0
                assert not app.state.document_service._processes
                monkeypatch.setattr(document_service, "_open_input_artifact_parent", original_open)
            if mode in {"root_revoke", "root_expire", "goal_pause", "goal_revision", "original_deadline", "late_cancel"}:
                from datetime import datetime, timedelta, timezone
                from src.db.models import Goal, OperatorSession
                original_parse = app.state.document_service.parse
                original_deadline = None
                async def change_after_positive_parse(*args, **kwargs):
                    nonlocal original_deadline
                    parsed = await original_parse(*args, **kwargs)
                    assert parsed["status"] == "succeeded" and parsed["cleanup"] == "wait_reaped"
                    async with factory.accounting_sessions() as db:
                        source_row = await db.get(WorkBoardInputArtifact, identifier)
                        stored = json.loads(source_row.document_metadata_json)
                        original_deadline = stored["execution_deadline"]
                        if mode.startswith("root_"):
                            root_row = await db.get(OperatorSession, source_row.owner_session_id)
                            if mode == "root_revoke": root_row.revoked_at = datetime.now(timezone.utc)
                            else: root_row.absolute_expires_at = datetime.now(timezone.utc)-timedelta(seconds=1)
                        elif mode.startswith("goal_"):
                            goal_row = await db.get(Goal, source_row.goal_id)
                            if mode == "goal_pause": goal_row.status = "paused"
                            else: goal_row.revision += 1
                        elif mode == "original_deadline":
                            from src.work_board import pipelines
                            expired_clock = datetime.fromisoformat(original_deadline)+timedelta(seconds=1)
                            class ExpiredClock(datetime):
                                @classmethod
                                def now(cls, tz=None): return expired_clock
                            monkeypatch.setattr(pipelines, "datetime", ExpiredClock)
                    if mode == "late_cancel": asyncio.current_task().cancel()
                    return parsed
                monkeypatch.setattr(app.state.document_service, "parse", change_after_positive_parse)
                if mode == "late_cancel":
                    from src.work_board.contracts import WorkBoardOwner
                    async def native_read():
                        async with factory.accounting_sessions() as db:
                            source_row = await db.get(WorkBoardInputArtifact, identifier)
                            native_owner = WorkBoardOwner(principal_id=source_row.owner_principal_id, session_id=source_row.owner_session_id)
                            return await app.state.document_service.read(db, native_owner, read_request(fmt, identifier))
                    cancelled_read = asyncio.create_task(native_read())
                    with pytest.raises(asyncio.CancelledError): await asyncio.wait_for(cancelled_read, 5)
                else:
                    refused = await client.post("/api/documents/read", json=read_request(fmt, identifier).model_dump())
                    assert refused.status_code in {401, 409}, refused.text
                async with factory.accounting_sessions() as db:
                    source_row = await db.get(WorkBoardInputArtifact, identifier)
                    retained_value = json.loads(source_row.document_metadata_json)
                    assert not retained_value.get("evidence")
                    assert source_row.document_reserved_bytes == 32*1024*1024
                    assert retained_value["live_writer"] is None
                    assert retained_value["execution_deadline"] == original_deadline
                    if mode == "late_cancel":
                        assert retained_value["reason"] == "document_output_cleanup_required"
                    else:
                        assert not list(root.glob(f"artifacts/work-board/document-sources/{identifier}/g*-evidence.fernet"))
                assert contacts == []
                return
            result = await client.post("/api/documents/read", json=read_request(fmt, identifier).model_dump())
            assert result.status_code == 200, result.text
            assert result.json()["status"] == "succeeded", result.text
            readback = await client.post("/api/documents/read", json=read_request(fmt, identifier).model_dump())
            assert readback.json() == result.json()
            listed = await client.get("/api/documents/sources")
            assert listed.json()["upload_readiness"] == "ready"
            assert [source["artifact_ref"] for source in listed.json()["sources"]] == [state["artifact_ref"]]
            assert contacts == []
            if mode == "csv":
                from datetime import datetime, timezone
                from src.db.models import OperatorSession
                from src.work_board.contracts import WorkBoardOwner
                from src.work_board.repository import BoardError
                async with factory.accounting_sessions() as db:
                    source_row = await db.get(WorkBoardInputArtifact, identifier)
                    owner = WorkBoardOwner(principal_id=source_row.owner_principal_id, session_id=source_row.owner_session_id)
                    original_root = await db.get(OperatorSession, owner.session_id)
                    original_root.revoked_at = datetime.now(timezone.utc)
                    await db.flush()
                    # The future typed native consumer cannot bypass current
                    # Root checks by omitting the authenticated HTTP operator.
                    with pytest.raises(BoardError) as revoked:
                        await app.state.document_service.read(db, owner, read_request(fmt, identifier))
                    assert revoked.value.code == "document_current_root_required"
                    await db.rollback()
            files = list(root.glob("artifacts/work-board/document-sources/**/*.fernet"))
            assert len(files) == 2 and all(file.stat().st_mode & 0o077 == 0 for file in files)
            assert all(raw not in file.read_bytes() for file in files)
            async with factory() as db:
                row = await db.get(WorkBoardInputArtifact, identifier)
                assert row.document_reserved_bytes == 32*1024*1024
                assert "Selected owner paragraph" not in row.document_metadata_json
            inspected = (await client.get(f"/api/documents/sources/{identifier}")).json()
            deleted = await client.delete(f"/api/documents/sources/{identifier}", params={"expected_revision": inspected["revision"]})
            assert deleted.status_code == 200, deleted.text
            assert deleted.json()["state"] == "deleted" and not list(root.glob("artifacts/work-board/document-sources/**/*.fernet"))
            if mode in {"restart", "cancel"}:
                await replacement.stop()
