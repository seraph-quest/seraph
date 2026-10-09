"""One real authenticated C1 build, private outputs and verified retirement."""
import json
from io import BytesIO

import httpx
import pytest
from fastapi import FastAPI
from sqlalchemy import select

from tests.test_inference_accounting import accounting_db
from tests.test_document_build_native_capacity import build_admission_lifecycle
from tests.test_general_task_planner import prepare, forbid_external_inference
from tests.test_document_build_storage import setup, SPEC


@pytest.mark.parametrize("kind", ["report", "table_workbook", "report_missing_pdf"])
async def test_actual_private_build_task_download_and_retire(accounting_db, monkeypatch, kind, build_admission_lifecycle):
    from src.api import documents, work_board
    from src.auth.service import authenticate_token
    from src.native_tools.registry import ToolRegistry
    from src.work_board.general_task import GeneralTaskService
    from src.work_board.dispatcher import WorkBoardDispatcher
    from src.db.models import WorkBoardInputArtifact, WorkBoardTask, WorkBoardAttempt, WorkflowRunState, WorkBoardEvent
    token, operator, owner, goal = await setup(accounting_db, monkeypatch)
    jobs, _owner = await prepare(accounting_db, monkeypatch, existing_owner=owner)
    await build_admission_lifecycle.start()
    sessions = accounting_db[2].accounting_sessions
    monkeypatch.setattr(documents, "get_session", sessions)
    registry = ToolRegistry(); registry.start()
    service = GeneralTaskService(registry); service.start()
    dispatcher = WorkBoardDispatcher(jobs=jobs, session_provider=sessions, general_tasks=service)
    monkeypatch.setattr(work_board, "dispatcher", dispatcher)
    app = FastAPI()
    @app.middleware("http")
    async def auth(request, call_next):
        request.state.operator = await authenticate_token(token, touch=False)
        return await call_next(request)
    app.include_router(documents.router, prefix="/api")
    app.include_router(work_board.router, prefix="/api")
    specification = dict(SPEC)
    if kind == "report_missing_pdf":
        specification = {**SPEC, "sections": [{**SPEC["sections"][0],
            "paragraphs": [SPEC["sections"][0]["paragraphs"][0] + " 🦉"]}]}
    if kind == "table_workbook":
        specification = {**SPEC, "kind": kind, "tables": [{"sheet_names": ["Results"],
            "cells": [{"sheet": "Results", "cell": "A1", "value": 2},
                {"sheet": "Results", "cell": "A2", "value": 3},
                {"sheet": "Results", "cell": "B1", "value": "=1+1"}],
            "formulas": [{"sheet": "Results", "cell": "A3", "expression": "SUM(A1:A2)"}], "formats": []}]}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://fixture") as client:
        created = await client.post("/api/documents/builds", json={"goal_id": goal.id,
            "goal_revision": 1, "spec": specification, "idempotency_key": "actual-" + kind})
        assert created.status_code == 200, created.text
        identifier = created.json()["build_id"]
        first_review = await client.get(f"/api/documents/builds/{identifier}/preview")
        assert first_review.status_code == 200, first_review.text
        request = {"expected_revision": first_review.json()["revision"],
            "review": first_review.json()["review"], "idempotency_key": "prepare-" + kind}
        proposed = await client.post(f"/api/documents/builds/{identifier}/prepare", json=request)
        assert proposed.status_code == 200, proposed.text
        task = proposed.json()["task"]
        task_id = task["task_id"]
        replay = await client.post(f"/api/documents/builds/{identifier}/prepare", json=request)
        assert replay.status_code == 200 and replay.json()["idempotent_replay"], replay.text
        denied = await client.post(f"/api/work-board/tasks/{task_id}/actions", json={
            "action": "promote", "expected_revision": task["task_revision"]})
        assert denied.status_code == 409, denied.text
        review = await client.get(f"/api/documents/builds/{identifier}/preview")
        assert review.status_code == 200 and review.json()["review"]["binding"]["task_id"] == task_id, review.text
        promoted = await client.post(f"/api/work-board/tasks/{task_id}/actions", json={
            "action": "promote", "expected_revision": task["task_revision"],
            "document_build_review": review.json()["review"]})
        assert promoted.status_code == 200, promoted.text
        outcome = await dispatcher.run_pass()
        if outcome["completed"] != 1:
            async with sessions() as db:
                runs = list((await db.scalars(select(WorkflowRunState))).all())
                print([{key: (await jobs.get_job(run.run_identity))[key]
                    for key in ("job_id", "status", "failure_reason", "result")} for run in runs])
        assert outcome["completed"] == 1, outcome
        listed = await client.get(f"/api/documents/builds/{identifier}/outputs")
        assert listed.status_code == 200, listed.text
        assert listed.json()["state"] == ("degraded" if kind == "report_missing_pdf" else "completed")
        editable = await client.get(f"/api/documents/builds/{identifier}/outputs/editable")
        pdf = await client.get(f"/api/documents/builds/{identifier}/outputs/pdf")
        assert editable.status_code == 200
        if kind == "report_missing_pdf":
            assert pdf.status_code == 409 and listed.json()["output"]["pdf_artifact"] is None
            assert "document_pdf_unsupported_glyph" in listed.json()["output"]["warnings"]
        else:
            assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
        for response in ((editable,) if kind == "report_missing_pdf" else (editable, pdf)):
            assert response.headers["cache-control"] == "no-store"
            assert response.headers["x-content-type-options"] == "nosniff"
            assert response.headers["content-disposition"].startswith("attachment;")
        if kind == "table_workbook":
            from openpyxl import load_workbook
            structure = load_workbook(BytesIO(editable.content), data_only=False)
            cached = load_workbook(BytesIO(editable.content), data_only=True)
            assert structure["Results"]["A3"].data_type == "f" and cached["Results"]["A3"].value == 5
            assert structure["Results"]["B1"].data_type == "s" and cached["Results"]["B1"].value == "=1+1"
        else:
            from docx import Document
            document = Document(BytesIO(editable.content))
            assert any("Private literal unavailable" in item.text for item in document.paragraphs)
        async with sessions() as db:
            build = await db.get(WorkBoardInputArtifact, identifier)
            current_task = await db.scalar(select(WorkBoardTask).where(WorkBoardTask.task_id == task_id))
            attempt = await db.scalar(select(WorkBoardAttempt).where(WorkBoardAttempt.task_id == task_id))
            runs = list((await db.scalars(select(WorkflowRunState))).all())
            events = list((await db.scalars(select(WorkBoardEvent).where(WorkBoardEvent.task_id == task_id))).all())
            assert len(runs) == 2
            child = next(run for run in runs if run.job_kind == "general_task_native_tool_v1")
            assert child.attempt_count == 1 and child.status == ("degraded" if kind == "report_missing_pdf" else "succeeded")
            step_receipt = next(item for item in json.loads(child.artifact_receipts_json)
                if item["artifact_type"] == "general_task_step")
            from src.work_board.dispatcher import _parse_typed_input
            assert all("Private literal unavailable" not in str(value) for value in [
                _parse_typed_input(current_task), current_task.title, current_task.body,
                *(event.metadata_json for event in events), *(run.arguments_json for run in runs),
                *(run.checkpoint_receipts_json for run in runs)])
            retirement = {"expected_revision": build.revision, "expected_task_revision": current_task.task_revision,
                "attempt_id": attempt.attempt_id, "idempotency_key": "retire-" + kind}
        from config.settings import settings
        from src.workspace import canonical_workspace_root
        original_step = canonical_workspace_root(settings.workspace_dir) / step_receipt["file_path"]
        original_bytes = original_step.read_bytes()
        original_step.write_bytes(original_bytes[:-1] + b" ")
        assert (await client.get(f"/api/documents/builds/{identifier}/outputs")).status_code == 409
        async with sessions() as db:
            assert (await db.get(WorkBoardInputArtifact, identifier)).document_reserved_bytes == 24*1024*1024
        original_step.write_bytes(original_bytes)
        premature = await client.request("DELETE", f"/api/documents/builds/{identifier}", json=retirement)
        assert premature.status_code == 409 and premature.json()["detail"]["code"] == "document_build_task_not_terminal"
        completed = await client.post(f"/api/work-board/tasks/{task_id}/actions", json={
            "action": "complete_review", "expected_revision": retirement["expected_task_revision"],
            "attempt_id": retirement["attempt_id"]})
        assert completed.status_code == 200, completed.text
        retirement["expected_task_revision"] = completed.json()["task"]["task_revision"]
        # Readback seal and SQL adoption do not make the filesystem atomic.
        # A later same-user replacement still denies download and retirement.
        from src.work_board import document_build_storage as storage, document_pairs as sources
        async with sessions() as db:
            private_row, private_value = await storage.owned(db, owner, identifier)
            private_editable = sources.source_path(private_row, private_value, "editable")
            original_cipher = private_editable.read_bytes()
        private_editable.write_bytes(original_cipher[:-1]+bytes([original_cipher[-1]^1]))
        assert (await client.get(f"/api/documents/builds/{identifier}/outputs/editable")).status_code == 409
        held_retirement = await client.request("DELETE", f"/api/documents/builds/{identifier}", json=retirement)
        assert held_retirement.status_code == 409, held_retirement.text
        async with sessions() as db:
            private_row, private_value = await storage.owned(db, owner, identifier)
            assert private_row.document_reserved_bytes == 24*1024*1024
            assert private_value["phase"] == "cleanup_tombstone"
            retirement["expected_revision"] = private_row.revision
        private_editable.write_bytes(original_cipher)
        retired = await client.request("DELETE", f"/api/documents/builds/{identifier}", json=retirement)
        assert retired.status_code == 200, retired.text
        assert retired.json()["state"] == "deleted" and retired.json()["quota_reserved_bytes"] == 0
        assert (await client.get(f"/api/documents/builds/{identifier}/outputs/editable")).status_code == 409
        snapshot = await jobs.inference_accounting_snapshot()
        assert snapshot["operation_count"] == 0
    service.stop(); registry.stop()
