from io import BytesIO
import json
from pathlib import Path
import struct
import subprocess
import sys
import pytest
from pydantic import ValidationError
from src.work_board.document_build_contracts import DocumentBuildSpec, SpreadsheetSpec, DocumentOutput, DOCX_MEDIA, PDF_MEDIA, canonical_spec
from src.work_board.document_build_formula import calculate, FormulaError
from src.work_board.document_build_renderer import render


def report(**changes):
    data = {'kind': 'report', 'title': 'Quarterly report', 'sections': [{'heading': 'Progress', 'paragraphs': ['Original local paragraph'], 'citation_refs': ['task:evidence:one']}], 'tables': [{'title': 'Results', 'columns': ['Item', 'Value'], 'rows': [['Work', 7]], 'citation_refs': []}], 'citations': [{'source_ref': 'task:evidence:one', 'label': 'Selected evidence'}], 'style_preset': 'plain'}
    return DocumentBuildSpec.model_validate(data | changes)


def spreadsheet(formulas=None, cells=None):
    return SpreadsheetSpec.model_validate({'sheet_names': ['Data'], 'cells': cells or [{'sheet': 'Data', 'cell': 'A1', 'value': 4}, {'sheet': 'Data', 'cell': 'A2', 'value': 6}], 'formulas': formulas or [], 'formats': []})


def workbook(sheet):
    return report(kind='table_workbook', tables=[sheet.model_dump()])


def test_docx_pdf_semantic_roundtrip():
    from docx import Document
    from pypdf import PdfReader
    result = render(report())
    document = Document(BytesIO(result.editable))
    assert 'Quarterly report' in [paragraph.text for paragraph in document.paragraphs]
    assert document.tables[0].cell(1, 1).text == '7'
    assert 'Selected evidence: task:evidence:one' in [paragraph.text for paragraph in document.paragraphs]
    reader = PdfReader(BytesIO(result.pdf))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert all(value in text for value in ('Quarterly report', 'Original local paragraph', 'Work', 'Selected evidence', 'task:evidence:one'))
    assert result.warnings == []


def test_xlsx_literal_prefix_and_formula_cache():
    from openpyxl import load_workbook
    literals = ['=1+1', '+SUM(A1:A2)', '-1+2', '@SUM(A1:A2)', 'https://example.invalid']
    cells = [{'sheet': 'Data', 'cell': f'A{index+1}', 'value': text} for index, text in enumerate(literals)]
    cells += [{'sheet': 'Data', 'cell': 'B1', 'value': 4}, {'sheet': 'Data', 'cell': 'B2', 'value': 6}]
    sheet = spreadsheet([{'sheet': 'Data', 'cell': 'B3', 'expression': '=SUM(B1:B2)'}], cells)
    result = render(workbook(sheet))
    for cached in (False, True):
        book = load_workbook(BytesIO(result.editable), data_only=cached)
        for index, text in enumerate(literals):
            cell = book['Data'][f'A{index+1}']
            assert cell.value == text and cell.data_type == 's'
            assert cell.hyperlink is None
        assert book['Data']['B3'].value == (10 if cached else '=SUM(B1:B2)')
    assert result.pdf is not None


@pytest.mark.parametrize('expression,expected', [('=SUM(A1:A2)', 10), ('=AVERAGE(A1:A2)', 5), ('=MIN(A1:A2)', 4), ('=MAX(A1:A2)', 6), ('=COUNT(A1:A2)', 2), ('=IF(A1<A2,3^2,0)', 9), ('=Data!$A$1+50%', 4.5), ('=-(A1+A2)/2', -5), ('=IF(TRUE,"exact text","other")', 'exact text')])
def test_finite_formulas(expression, expected):
    sheet = spreadsheet([{'sheet': 'Data', 'cell': 'B1', 'expression': expression}])
    assert calculate(sheet)[('Data', 0, 1)] == expected


@pytest.mark.parametrize('expression', ['=1/0', '=WEBSERVICE("https://bad")', '=IF(FALSE,WEBSERVICE("bad"),1)', '=NOW()', '=SUM([Book]Sheet!A1)', '=cmd|"/C calc"!A1', '=Other!A1', '=A257', '=BM1', '=2^1000000', '=TRUE+1', '=IF(1,2,3)', '=AVERAGE(C1:C2)', '=A1:A2+1', '=SUM(A1:A256,A1:BL256)', '=1e999', '=__import__("os")'])
def test_unsupported_formula_denied(expression):
    sheet = spreadsheet([{'sheet': 'Data', 'cell': 'B1', 'expression': expression}])
    with pytest.raises(ValueError):
        calculate(sheet)


def test_cycles_hidden_if_and_overlap_denied():
    sheet = spreadsheet([{'sheet': 'Data', 'cell': 'B1', 'expression': '=IF(FALSE,B2,1)'}, {'sheet': 'Data', 'cell': 'B2', 'expression': '=B1'}])
    with pytest.raises(FormulaError, match='cycle'):
        calculate(sheet)
    with pytest.raises(ValidationError):
        spreadsheet([{'sheet': 'Data', 'cell': 'A1', 'expression': '=2'}])


@pytest.mark.parametrize('change', [{'extra': 1}, {'title': 'a'*201}, {'sections': [{'heading': 'h', 'paragraphs': ['é'*3000], 'citation_refs': []}]}, {'citations': [{'source_ref': 'x', 'label': 'l', 'extra': True}]}, {'tables': [{'title': 't', 'columns': ['a'], 'rows': [[1, 2]], 'citation_refs': []}]}])
def test_closed_bounds(change):
    with pytest.raises(ValidationError):
        report(**change)


def test_missing_pdf_keeps_editable_and_warning():
    from docx import Document
    result = render(report(title='Unsupported glyph 😀'))
    assert result.pdf is None
    assert result.warnings == ['document_pdf_unsupported_glyph']
    assert Document(BytesIO(result.editable)).paragraphs[0].text == 'Unsupported glyph 😀'


def test_missing_pdf_module_preserves_editable(monkeypatch):
    import src.work_board.document_build_renderer as renderer
    def absent(*args):
        raise ImportError('private content must not reach warning')
    monkeypatch.setattr(renderer, 'pdf', absent)
    result = renderer.render(report())
    assert result.editable.startswith(b'PK') and result.pdf is None
    assert result.warnings == ['document_pdf_render_unavailable']


def test_formula_rejected_at_spec_boundary():
    with pytest.raises(ValidationError, match='Data!B1') as caught:
        workbook(spreadsheet([{'sheet': 'Data', 'cell': 'B1', 'expression': '=1/0'}]))
    error = caught.value.errors()[0]
    assert error['type'] == 'document_formula_invalid'
    assert error['ctx'] == {'sheet': 'Data', 'cell': 'B1', 'formula_code': 'document_formula_arithmetic_invalid'}


def test_formula_error_context_never_contains_expression():
    with pytest.raises(ValidationError) as caught:
        workbook(spreadsheet([{'sheet': 'Data', 'cell': 'B1', 'expression': '=PRIVATE_SECRET(1)'}]))
    error = caught.value.errors()[0]
    assert error['ctx'] == {'sheet': 'Data', 'cell': 'B1', 'formula_code': 'document_formula_function_unsupported'}
    assert 'PRIVATE_SECRET' not in error['msg']


def test_blank_cell_format_and_boolean_formula_cache():
    from openpyxl import load_workbook
    sheet = spreadsheet([{'sheet': 'Data', 'cell': 'B1', 'expression': '=A1<A2'}])
    data = sheet.model_dump()
    data['formats'] = [{'sheet': 'Data', 'cell': 'D5', 'preset': 'decimal'}]
    result = render(workbook(SpreadsheetSpec.model_validate(data)))
    book = load_workbook(BytesIO(result.editable), data_only=True)
    assert book['Data']['B1'].value is True
    assert book['Data']['D5'].number_format == '0.00'


@pytest.mark.parametrize('change', [
    {'sheet_names': ['Data', 'data']}, {'sheet_names': ['bad/sheet']},
    {'formulas': [{'sheet': 'Data', 'cell': 'B1', 'expression': '=1', 'cached_value': 1}]},
    {'cells': [{'sheet': 'Data', 'cell': 'A1', 'value': float('nan')}]},
    {'cells': [{'sheet': 'Data', 'cell': 'A1', 'value': '\x00'}]},
    {'cells': [{'sheet': 'Data', 'cell': 'A1', 'value': 1}] * 2},
    {'formats': [{'sheet': 'Data', 'cell': 'A1', 'preset': 'arbitrary'}]},
])
def test_nested_spreadsheet_denials(change):
    with pytest.raises(ValidationError):
        SpreadsheetSpec.model_validate(spreadsheet().model_dump() | change)


def test_aggregate_spec_bytes_and_text_control_denied():
    with pytest.raises(ValidationError, match='spec_byte_bound'):
        report(sections=[{'heading': 'h', 'paragraphs': ['a'*4096]*17, 'citation_refs': []}])
    with pytest.raises(ValidationError, match='character_invalid'):
        report(title='bad\x00')


def test_closed_output_four_fields_and_missing_pdf():
    artifact = {'artifact_ref': 'document-build:00000000-0000-0000-0000-000000000001:editable', 'sha256': 'a'*64, 'size_bytes': 1, 'media_type': DOCX_MEDIA}
    value = {'editable_artifact': artifact, 'pdf_artifact': None, 'source_refs': [], 'warnings': ['document_pdf_render_unavailable']}
    assert set(DocumentOutput.model_validate(value).model_dump()) == set(value)
    for change in ({'warnings': []}, {'extra': True}, {'source_refs': ['one', 'one']}, {'pdf_artifact': artifact | {'artifact_ref': 'document-build:00000000-0000-0000-0000-000000000002:pdf', 'media_type': PDF_MEDIA}}):
        with pytest.raises(ValidationError):
            DocumentOutput.model_validate(value | change)


@pytest.mark.parametrize('wire', [struct.pack('!II', 65537, 0), struct.pack('!II', 1, 16385), struct.pack('!II', 1, 0)+b'x'+b'extra'])
def test_actual_child_pipe_bounds_closed_reaped(wire):
    child = Path(__file__).parents[1] / 'src/work_board/document_build_child.py'
    process = subprocess.Popen([sys.executable, '-I', '-B', str(child), 'b'*32], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env={}, close_fds=True)
    try:
        assert json.loads(process.stdout.readline())['state'] == 'ready'
        output, error = process.communicate(wire, timeout=30)
        assert process.wait(timeout=5) == 0, error
        editable, pdf, size = struct.unpack('!III', output[:12])
        assert editable == pdf == 0 and len(output) == 12+size
        assert json.loads(output[12:])['status'] == 'blocked'
    finally:
        if process.poll() is None:
            process.kill()
            process.wait(timeout=5)
        for stream in (process.stdin, process.stdout, process.stderr):
            stream.close()


def test_actual_confined_child_roundtrip_positive_reap():
    child = Path(__file__).parents[1] / 'src/work_board/document_build_child.py'
    process = subprocess.Popen([sys.executable, '-I', '-B', str(child), 'a'*32], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env={}, close_fds=True)
    try:
        ready = json.loads(process.stdout.readline())
        assert ready == {'state': 'ready', 'nonce': 'a'*32, 'network_denied': True, 'profile': 'document-build-renderer.v1'}
        raw = canonical_spec(report())
        output, error = process.communicate(struct.pack('!II', len(raw), 0) + raw, timeout=30)
        assert process.returncode == 0, error
        editable_size, pdf_size, metadata_size = struct.unpack('!III', output[:12])
        assert len(output) == 12 + editable_size + pdf_size + metadata_size
        assert editable_size > 0 and pdf_size > 0
        metadata = json.loads(output[-metadata_size:])
        assert metadata['status'] == 'succeeded' and metadata['provider_contacts'] == 0
        assert process.wait(timeout=5) == 0
    finally:
        if process.poll() is None:
            process.kill()
            process.wait(timeout=5)
        for stream in (process.stdin, process.stdout, process.stderr):
            stream.close()


@pytest.mark.parametrize('entry', ['calculate', 'spec_validation'])
def test_range_budget_denies_second_full_expansion_before_iterator(entry, monkeypatch):
    import builtins
    import src.work_board.document_build_formula as formulas
    sheet = SpreadsheetSpec.model_validate({'sheet_names': ['Data', 'Output'],
        'cells': [{'sheet': 'Data', 'cell': 'A1', 'value': 4}],
        'formulas': [{'sheet': 'Output', 'cell': 'A1', 'expression': '=SUM(Data!A1:BL256,Data!A1:BL256)'}],
        'formats': []})
    entries = []
    def bounded_range(*args):
        if args == (0, 256):
            entries.append(args)
            if len(entries) > 1:
                raise AssertionError('over-budget full range began allocation')
        return builtins.range(*args)
    monkeypatch.setattr(formulas, 'range', bounded_range, raising=False)
    if entry == 'calculate':
        with pytest.raises(FormulaError) as caught:
            calculate(sheet)
        assert (caught.value.code, caught.value.sheet, caught.value.cell) == ('document_formula_range_bound', 'Output', 'A1')
    else:
        with pytest.raises(ValidationError) as caught:
            workbook(sheet)
        assert caught.value.errors()[0]['ctx'] == {'sheet': 'Output', 'cell': 'A1', 'formula_code': 'document_formula_range_bound'}
    assert len(entries) == 1


def test_range_budget_is_shared_across_formulas_before_one_cell_allocation(monkeypatch):
    import builtins
    import src.work_board.document_build_formula as formulas
    sheet = SpreadsheetSpec.model_validate({'sheet_names': ['Data', 'Output'],
        'cells': [{'sheet': 'Data', 'cell': 'A1', 'value': 4}],
        'formulas': [{'sheet': 'Output', 'cell': 'A1', 'expression': '=SUM(Data!A1:BL256)'},
            {'sheet': 'Output', 'cell': 'A2', 'expression': '=SUM(Data!A1:A1)'}], 'formats': []})
    def bounded_range(*args):
        if args == (0, 1):
            raise AssertionError('over-budget one-cell range began allocation')
        return builtins.range(*args)
    monkeypatch.setattr(formulas, 'range', bounded_range, raising=False)
    with pytest.raises(FormulaError) as caught:
        calculate(sheet)
    assert (caught.value.code, caught.value.sheet, caught.value.cell) == ('document_formula_range_bound', 'Output', 'A2')


@pytest.mark.parametrize('split_formulas', [False, True])
def test_exact_total_range_budget_preserves_valid_calculation(split_formulas):
    expressions = ['=SUM(Data!A1:AF256)', '=SUM(Data!AG1:BL256)'] if split_formulas else ['=SUM(Data!A1:AF256,Data!AG1:BL256)']
    sheet = SpreadsheetSpec.model_validate({'sheet_names': ['Data', 'Output'],
        'cells': [{'sheet': 'Data', 'cell': 'A1', 'value': 4}, {'sheet': 'Data', 'cell': 'AG1', 'value': 6}],
        'formulas': [{'sheet': 'Output', 'cell': f'A{index+1}', 'expression': expression}
            for index, expression in enumerate(expressions)], 'formats': []})
    values = calculate(sheet)
    assert values[('Output', 0, 0)] == (4 if split_formulas else 10)
    if split_formulas:
        assert values[('Output', 1, 0)] == 6
    assert workbook(sheet).kind == 'table_workbook'
