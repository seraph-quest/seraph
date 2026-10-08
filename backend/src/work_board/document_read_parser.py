"""Fixed literal parsers. Called only after child confinement, never inference."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import PurePosixPath
import zipfile

MIB = 1024 * 1024
SOURCE_LIMIT = 16 * MIB
OUTPUT_LIMIT = MIB
CELL_LIMIT = 100_000


class DocumentReadError(ValueError):
    pass


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def spreadsheet_formula_literal(value):
    """Formula objects are metadata, never evidence through object repr."""
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    if isinstance(value, DataTableFormula):
        # This object has input/range attributes but no source expression text.
        # Converting attributes into an expression would invent a formula.
        raise DocumentReadError("document_spreadsheet_data_table_formula_unsupported")
    if isinstance(value, ArrayFormula):
        value = value.text
    if not isinstance(value, str) or not value.startswith("=") or len(value) < 2:
        raise DocumentReadError("document_spreadsheet_formula_literal_unsupported")
    if len(value.encode("utf-8")) > OUTPUT_LIMIT:
        raise DocumentReadError("document_output_size_exceeded")
    return value


def office_package(raw):
    """Preflight central directory before libraries expand XML; extract nothing."""
    try:
        package = zipfile.ZipFile(io.BytesIO(raw))
        members = package.infolist()
        if len(members) > 2048 or sum(item.file_size for item in members) > 32 * MIB:
            raise DocumentReadError("document_zip_expansion_exceeded")
        names = set()
        for item in members:
            path = PurePosixPath(item.filename)
            if (item.filename in names or path.is_absolute() or ".." in path.parts
                or "\\" in item.filename or ":" in item.filename
                or (item.external_attr >> 16) & 0o170000 == 0o120000):
                raise DocumentReadError("document_zip_path_unsupported")
            names.add(item.filename)
            if (item.flag_bits & 1 or item.file_size > 16 * MIB
                or item.file_size > max(1, item.compress_size) * 100):
                raise DocumentReadError("document_zip_expansion_or_protection_unsupported")
            lower = item.filename.lower()
            if "vbaproject" in lower or lower.startswith("xl/externallinks/"):
                raise DocumentReadError("document_macros_or_external_links_unsupported")
        # External relationships are never dereferenced. Reject instead of
        # silently implying that linked material was incorporated.
        for item in members:
            if item.filename.endswith((".rels", ".xml")):
                contents = package.read(item)
                if b"<!DOCTYPE" in contents.upper() or b"<!ENTITY" in contents.upper():
                    raise DocumentReadError("document_xml_entities_unsupported")
                if item.filename == "[Content_Types].xml" and b"macroEnabled" in contents:
                    raise DocumentReadError("document_macros_or_external_links_unsupported")
                if item.filename.endswith(".rels"):
                    from xml.etree import ElementTree
                    relationships = ElementTree.fromstring(contents)
                    if any(element.get("TargetMode", "").lower() == "external" for element in relationships.iter()):
                        raise DocumentReadError("document_external_relationship_unsupported")
        package.close()
    except DocumentReadError:
        raise
    except Exception:
        raise DocumentReadError("document_office_malformed_or_protected") from None


def extract(raw, request):
    if not 0 < len(raw) <= SOURCE_LIMIT:
        raise DocumentReadError("document_source_size_exceeded")
    fmt = request["format"]
    selection = request["selection"]
    limits = request["page_sheet_limits"]
    ref = request["artifact_ref"]
    sections, warnings = [], []
    output_bytes = 0
    def add(location, text="", cells=None):
        nonlocal output_bytes
        section = {"source_ref": ref + "#" + location, "text": text, "table_cells": cells or []}
        output_bytes += len(canonical(section)) + 1
        if output_bytes > OUTPUT_LIMIT - 8192:
            raise DocumentReadError("document_output_size_exceeded")
        sections.append(section)
    if fmt in {"docx", "xlsx"}:
        if fmt == "docx" and len(raw) > 10 * MIB:
            raise DocumentReadError("document_docx_size_exceeded")
        office_package(raw)
    try:
        if fmt == "pdf":
            import pypdf
            if pypdf.__version__ != "6.19.0":
                raise DocumentReadError("document_parser_version_blocked")
            from pypdf import PdfReader, apply_configuration
            with apply_configuration(zlib_maximum_output_length=16*MIB,
                array_based_stream_maximum_output_length=16*MIB,
                maximum_declared_stream_length=16*MIB, page_tree_maximum_entries=101,
                xform_maximum_invocations_per_extraction=100,
                image_maximum_buffer_size=16*MIB, jbig2dec_binary=None):
                reader = PdfReader(io.BytesIO(raw), strict=True)
                if reader.is_encrypted:
                    raise DocumentReadError("document_pdf_protected_unsupported")
                if len(reader.pages) > 100:
                    raise DocumentReadError("document_pdf_page_limit_exceeded")
                selected = selection["pages"] or list(range(1, len(reader.pages)+1))
                if any(page > len(reader.pages) for page in selected):
                    raise DocumentReadError("document_page_selection_unavailable")
                selected = selected[:limits["max_pages"]]
                skipped = sorted(set(range(1, len(reader.pages)+1)) - set(selected))
                if skipped:
                    warnings.append("Skipped physical PDF pages: " + ",".join(map(str, skipped)))
                for page in selected:
                    text = reader.pages[page-1].extract_text()
                    if not text or not text.strip():
                        raise DocumentReadError("document_pdf_scanned_or_empty_unsupported")
                    add(f"page={page}", text)
        elif fmt == "docx":
            import docx
            if docx.__version__ != "1.2.0":
                raise DocumentReadError("document_parser_version_blocked")
            from docx import Document
            document = Document(io.BytesIO(raw))
            for ordinal, paragraph in enumerate(document.paragraphs, 1):
                if paragraph.text:
                    add(f"paragraph={ordinal}", paragraph.text)
            count = 0
            for table_index, table in enumerate(document.tables, 1):
                for row_index, row in enumerate(table.rows, 1):
                    cells = []
                    for column, cell in enumerate(row.cells, 1):
                        if cell.text:
                            count += 1
                            if count > limits["max_cells"]:
                                raise DocumentReadError("document_populated_cell_limit_exceeded")
                            cells.append({"source_ref": ref+f"#table={table_index}&row={row_index}&column={column}",
                                "text": cell.text, "formula": None, "cached_value": None})
                    if cells:
                        add(f"table={table_index}&row={row_index}", cells=cells)
        elif fmt == "xlsx":
            import openpyxl
            if openpyxl.__version__ != "3.1.5":
                raise DocumentReadError("document_parser_version_blocked")
            from openpyxl import load_workbook
            from urllib.parse import quote
            literal = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            cached = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            try:
                selected = selection["sheets"] or literal.sheetnames
                if any(sheet not in literal.sheetnames for sheet in selected):
                    raise DocumentReadError("document_sheet_selection_unavailable")
                selected = selected[:limits["max_sheets"]]
                skipped = [sheet for sheet in literal.sheetnames if sheet not in selected]
                if skipped:
                    warnings.append("Skipped XLSX sheets: " + ",".join(skipped)[:2048])
                count = 0
                for name in selected:
                    sheet = literal[name]
                    # Bound sparse coordinate traversal as well as populated cells.
                    if sheet.max_row is None or sheet.max_column is None or sheet.max_row*sheet.max_column > CELL_LIMIT:
                        raise DocumentReadError("document_sheet_coordinate_bound_exceeded")
                    for row, values in zip(sheet.iter_rows(), cached[name].iter_rows()):
                        cells = []
                        for cell, value in zip(row, values):
                            if cell.value is None:
                                continue
                            count += 1
                            if count > limits["max_cells"]:
                                raise DocumentReadError("document_populated_cell_limit_exceeded")
                            formula = spreadsheet_formula_literal(cell.value) if cell.data_type == "f" else None
                            cells.append({"source_ref": ref+f"#sheet={quote(name, safe='')}&cell={cell.coordinate}",
                                "text": formula if formula is not None else str(cell.value), "formula": formula,
                                "cached_value": str(value.value) if formula and value.value is not None else None})
                        if cells:
                            add(f"sheet={quote(name, safe='')}&row={row[0].row}", cells=cells)
                warnings.append("Formula literals were not evaluated; cached value freshness is unknown.")
            finally:
                literal.close(); cached.close()
        elif fmt == "csv":
            def get_column_letter(column):
                name = ""
                while column:
                    column, remainder = divmod(column-1, 26)
                    name = chr(65+remainder) + name
                return name
            try:
                text = raw.decode("utf-8-sig", errors="strict")
            except UnicodeError:
                raise DocumentReadError("document_csv_utf8_required") from None
            csv.field_size_limit(OUTPUT_LIMIT)
            count = 0
            for row_number, row in enumerate(csv.reader(io.StringIO(text, newline=""), strict=True), 1):
                if row_number > CELL_LIMIT or len(row) > CELL_LIMIT:
                    raise DocumentReadError("document_sheet_coordinate_bound_exceeded")
                cells = []
                for column, value in enumerate(row, 1):
                    if value:
                        count += 1
                        if count > limits["max_cells"]:
                            raise DocumentReadError("document_populated_cell_limit_exceeded")
                        cells.append({"source_ref": ref+f"#sheet=CSV&cell={get_column_letter(column)}{row_number}",
                            "text": value, "formula": value if value.startswith("=") else None, "cached_value": None})
                if cells:
                    add(f"sheet=CSV&row={row_number}", cells=cells)
        else:
            raise DocumentReadError("document_format_unsupported")
    except DocumentReadError:
        raise
    except ImportError:
        raise DocumentReadError("document_parser_dependency_unavailable") from None
    except Exception:
        raise DocumentReadError("document_malformed_or_unsupported") from None
    if not sections:
        raise DocumentReadError("document_empty_unsupported")
    result = {"sections": sections, "warnings": warnings,
        "source_digest": hashlib.sha256(raw).hexdigest(), "no_learning": True}
    if len(canonical(result)) > OUTPUT_LIMIT:
        raise DocumentReadError("document_output_size_exceeded")
    return result
